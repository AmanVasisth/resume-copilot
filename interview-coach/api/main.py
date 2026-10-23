import os, json, re
from io import BytesIO
from pathlib import Path
from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.responses import StreamingResponse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from fastapi.middleware.cors import CORSMiddleware
from .models import InterviewRequest, InterviewPlan, EvaluateAnswerRequest, AnswerEvaluation, NextQuestionRequest, NextQuestionResponse, Question
from .engine import build_profile, make_questions, evaluate_answer
from .llm import evaluate_with_llm, rewrite_resume_bullet
from .scoring import apply_rubric_result
from .adaptive import InterviewState, select_next_question, readiness_report
from .resume_parser import extract_resume_text
from . import database

app = FastAPI(title="CareerCoach AI Interview Engine", version="0.4.3")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("startup")
def startup():
    database.init_db()

@app.get("/health")
def health():
    return {"status":"ok","service":"careercoach-interview-engine","version":"0.4.3"}

@app.post("/v1/resume/extract")
async def extract_resume(file: UploadFile = File(...)):
    try:
        content = await file.read()
        text = extract_resume_text(file.filename or "", content)
        if not text:
            raise ValueError("No readable text found in the uploaded resume.")
        return {"filename": file.filename, "text": text}
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc))

@app.post("/v1/resume/export")
async def export_revised_resume(file: UploadFile = File(...), approved_bullets: str = ""):
    """Export approved edits to a polished, editable DOCX resume."""
    try:
        content = await file.read()
        ext = Path(file.filename or "").suffix.lower()
        replacements = json.loads(approved_bullets or "[]")
        if not isinstance(replacements, list):
            raise ValueError("Approved bullet changes must be a list.")

        if ext == ".docx":
            doc = Document(BytesIO(content))
            for item in replacements:
                original = str(item.get("original", "")).strip()
                replacement = str(item.get("replacement", "")).strip()
                if not original or not replacement:
                    continue
                for paragraph in doc.paragraphs:
                    if paragraph.text.strip() == original:
                        if paragraph.runs:
                            paragraph.runs[0].text = replacement
                            for run in paragraph.runs[1:]:
                                run.text = ""
                        else:
                            paragraph.add_run(replacement)
                        break
        else:
            resume_text = extract_resume_text(file.filename or "", content)
            if not resume_text.strip():
                raise ValueError("No readable resume text found.")
            for item in replacements:
                original = str(item.get("original", "")).strip()
                replacement = str(item.get("replacement", "")).strip()
                if original and replacement:
                    resume_text = resume_text.replace(original, replacement, 1)

            # Repair frequent PDF text-extraction artefacts without adding claims.
            resume_text = re.sub(r"(?i)g\\s+mail\\.com", "gmail.com", resume_text)
            resume_text = re.sub(r"(?i)Microsof\\s+t", "Microsoft", resume_text)
            for broken, fixed in {
                "Develop ed": "Developed", "develop ed": "developed",
                "cont rolled": "controlled", "app lication": "application",
                "reason ing": "reasoning", "docume nt": "document",
                "groundn ess": "groundedness", "deploy ment": "deployment",
                "sou rce": "source", "ve ctor": "vector", "Microsof t": "Microsoft",
                "DevO ps": "DevOps", "wit h": "with", "Po wer BI": "Power BI",
            }.items():
                resume_text = resume_text.replace(broken, fixed)
            resume_text = re.sub(r"\\[X\\]\\s*%\\s*improvement\\s+in\\s+\\[metric\\]", "", resume_text, flags=re.I)
            resume_text = re.sub(r"\\s+([,.;:])", r"\\1", resume_text)
            # Build a consistent, recruiter-friendly DOCX instead of dumping every
            # extracted PDF line into unformatted paragraphs. Text is not invented.
            doc = Document()
            section = doc.sections[0]
            section.top_margin = Inches(0.48)
            section.bottom_margin = Inches(0.48)
            section.left_margin = Inches(0.62)
            section.right_margin = Inches(0.62)

            styles = doc.styles
            normal = styles["Normal"]
            normal.font.name = "Arial"
            normal.font.size = Pt(9)
            normal.font.color.rgb = RGBColor(38, 48, 65)
            normal.paragraph_format.space_after = Pt(2)
            normal.paragraph_format.line_spacing = 1.05

            heading_names = {
                "professional summary", "profile summary", "summary",
                "experience", "professional experience", "work experience",
                "projects", "key projects", "technical skills", "skills",
                "core competencies", "education", "certifications", "awards",
                "achievements", "key result areas"
            }
            lines = [re.sub(r"^\s*[•▪●]\s*", "", line).strip()
                     for line in resume_text.splitlines()]
            lines = [line for line in lines if line]
            if not lines:
                raise ValueError("No readable resume text found.")

            # Name and contact line are highlighted, while the original wording
            # and sequence are retained as far as text extraction allows.
            title = doc.add_paragraph()
            title.paragraph_format.space_after = Pt(2)
            title_run = title.add_run(lines[0])
            title_run.bold = True
            title_run.font.name = "Arial"
            title_run.font.size = Pt(19)
            title_run.font.color.rgb = RGBColor(37, 49, 78)

            for line in lines[1:]:
                if line.lower() in heading_names:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(7)
                    p.paragraph_format.space_after = Pt(3)
                    run = p.add_run(line.upper())
                    run.bold = True
                    run.font.name = "Arial"
                    run.font.size = Pt(10.5)
                    run.font.color.rgb = RGBColor(79, 70, 229)
                elif re.match(r"^(Qlik Sense Developer|Senior Consultant|Consultant|Power BI Developer|GenAI|AI Engineer|Data Analyst).{0,180}\|\s*", line, re.I):
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run(line)
                    run.bold = True
                    run.font.color.rgb = RGBColor(37, 49, 78)
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.keep_together = True
                    if line.startswith(("-", "•", "▪", "●")):
                        p.paragraph_format.left_indent = Inches(0.16)
                        p.add_run("• " + line.lstrip("-•▪● "))
                    else:
                        p.add_run(line)

        stream = BytesIO()
        doc.save(stream)
        stream.seek(0)
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": 'attachment; filename="careercoach-revised-resume.docx"'},
        )
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not export revised resume: {exc}")

@app.post("/v1/resume/rewrite-bullet")
def rewrite_bullet(payload: dict):
    original = str(payload.get("original", "")).strip()
    draft = str(payload.get("draft", "")).strip()
    job_description = str(payload.get("job_description", "")).strip()
    if not original:
        raise HTTPException(status_code=400, detail="A resume bullet is required.")
    if len(original) > 3000 or len(draft) > 3000 or len(job_description) > 12000:
        raise HTTPException(status_code=413, detail="Please shorten the bullet or job description and try again.")
    result = rewrite_resume_bullet(original, draft, job_description)
    if not result:
        raise HTTPException(status_code=503, detail="AI rewriting is temporarily unavailable. Check the configured AI provider and retry.")
    return result

@app.post("/v1/interview/plan", response_model=InterviewPlan)
def create_plan(req: InterviewRequest):
    profile = build_profile(req)
    memory_topics = []
    if req.candidate_key:
        try:
            memory_topics = database.detailed_recommendations(req.candidate_key).get("topics", [])
        except Exception:
            memory_topics = []
    questions = make_questions(profile, req.question_count, req.mode, memory_topics=memory_topics)
    return InterviewPlan(profile=profile, questions=questions)

@app.post("/v1/sessions")
def start_session(payload: dict):
    key = str(payload.get("candidate_key", "")).strip()
    role = str(payload.get("target_role", "")).strip()
    if not key or not role:
        raise HTTPException(status_code=400, detail="candidate_key and target_role are required.")
    try:
        return database.create_session(key, role, float(payload.get("years_experience", 0)), str(payload.get("mode", "mixed")), payload.get("plan"))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc))

@app.get("/v1/sessions")
def sessions(candidate_key: str, limit: int = 25):
    return {"sessions": database.list_sessions(candidate_key, limit)}

@app.get("/v1/progress")
def progress(candidate_key: str):
    return database.progress_summary(candidate_key)

@app.get("/v1/recommendations")
def recommendations(candidate_key: str):
    return database.detailed_recommendations(candidate_key)

@app.get("/v1/sessions/{session_id}")
def session_detail(session_id: int, candidate_key: str):
    try:
        return database.get_session(session_id, candidate_key)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc))

@app.post("/v1/sessions/{session_id}/answers")
def save_session_answer(session_id: int, payload: dict):
    try:
        return database.save_answer(
            session_id, str(payload.get("candidate_key", "")),
            payload["question"], str(payload.get("answer", "")), payload["evaluation"]
        )
    except (ValueError, KeyError) as exc:
        raise HTTPException(status_code=400, detail=str(exc))

@app.post("/v1/sessions/{session_id}/complete")
def complete_session(session_id: int, payload: dict):
    try:
        return database.complete_session(session_id, str(payload.get("candidate_key", "")), int(payload["overall_score"]))
    except (ValueError, KeyError) as exc:
        raise HTTPException(status_code=400, detail=str(exc))

@app.post("/v1/interview/evaluate", response_model=AnswerEvaluation)
def evaluate(req: EvaluateAnswerRequest):
    normalized = " ".join(req.answer.lower().split())
    short_or_nonanswer = len(req.answer.strip().split()) < 3 or normalized in {"no", "nope", "n/a", "na", "idk", "dont know", "not sure", "skip", "pass", "nothing"}
    if short_or_nonanswer:
        result = evaluate_answer(req.question, req.answer, req.candidate)
    else:
        llm_result = evaluate_with_llm(req.question, req.answer, req.candidate)
        if not llm_result:
            raise HTTPException(
                status_code=503,
                detail="AI evaluation is temporarily unavailable. Your answer has not been scored; please retry shortly.",
            )
        try:
            provider_used = llm_result.get(
                "_evaluation_provider",
                "openrouter" if os.getenv("LLM_PROVIDER") == "openrouter" else "gemini",
            )
            result = apply_rubric_result(llm_result, provider_used)
        except Exception:
            raise HTTPException(
                status_code=503,
                detail="AI evaluation returned an invalid response. Your answer has not been scored; please retry.",
            )
    # Adaptive probing is only needed for a weak answer, and only one level deep.
    if req.question.id.startswith("followup-") or result.score >= 65:
        result.follow_up_question = None
    return result

@app.post("/v1/interview/next", response_model=NextQuestionResponse)
def next_question(req: NextQuestionRequest):
    state=InterviewState(asked=req.asked_questions, evaluations=req.evaluations)
    if req.last_question.id not in {q.id for q in state.asked}:
        state.record(req.last_question, req.last_evaluation)
    if (
        req.last_evaluation.score < 65
        and req.last_evaluation.follow_up_question
        and not req.last_question.id.startswith("followup-")
    ):
        follow_up = Question(
            id=f"followup-{len(req.asked_questions) + 1}",
            category=req.last_question.category,
            skill=req.last_question.skill,
            difficulty=req.last_question.difficulty,
            question=req.last_evaluation.follow_up_question,
            rationale="Adaptive probe generated from the candidate's previous answer.",
            source="adaptive",
        )
        return NextQuestionResponse(
            question=follow_up,
            interview_complete=False,
            readiness=readiness_report(req.candidate, state),
        )

    q=select_next_question(req.candidate, req.available_questions, state, req.last_question, req.last_evaluation)
    return NextQuestionResponse(question=q, interview_complete=q is None, readiness=readiness_report(req.candidate, state))
